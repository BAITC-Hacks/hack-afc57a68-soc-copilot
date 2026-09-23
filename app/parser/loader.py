"""Загрузка PDF / DOCX / XLSX в модель Document с пунктами."""
from __future__ import annotations

import re
from pathlib import Path

from app.models import Clause, Document
from app.parser.clauses import split_into_clauses


def _meta_from_text(first_page: str) -> dict:
    t = re.sub(r"\s+", " ", first_page)
    ed = re.search(r"редакция\s*(?:No|№|N)?\s*(\d+)", t, re.I)
    prot = re.search(r"Протокол\s*(?:No|№|N)?\s*(\d+)\s*от\s*«?(\d{1,2})»?\s*(\w+)\s*(\d{4})", t, re.I)
    title = re.search(r"(ПОЛОЖЕНИЕ\s+О\s+[А-ЯЁ\s«»\"]+?)(?:\(|АО|$)", t)
    return {
        "edition": ed.group(1) if ed else "",
        "approved": (f"Протокол № {prot.group(1)} от {prot.group(2)} {prot.group(3)} {prot.group(4)}"
                     if prot else ""),
        "title": title.group(1).strip().capitalize() if title else "",
    }


def load_pdf(path: str, doc_id: str) -> tuple[Document, list[dict]]:
    import pymupdf  # PyMuPDF

    pdf = pymupdf.open(path)
    pages = [(i + 1, pdf[i].get_text()) for i in range(len(pdf))]
    meta = _meta_from_text(pages[0][1] if pages else "")
    clauses, anomalies = split_into_clauses(pages, doc_id, meta["edition"])
    doc = Document(doc_id=doc_id, file=Path(path).name, kind="pdf", clauses=clauses, **meta)
    return doc, anomalies


def _docx_paragraph_texts(path: str) -> list[str]:
    """Возвращает абзацы DOCX, восстанавливая автонумерацию списков (w:numPr)."""
    import docx

    d = docx.Document(path)
    counters: dict[tuple[str, int], int] = {}
    out: list[str] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        numpr = p._p.pPr.numPr if p._p.pPr is not None else None
        if numpr is not None and numpr.numId is not None:
            num_id = str(numpr.numId.val)
            lvl = int(numpr.ilvl.val) if numpr.ilvl is not None else 0
            counters[(num_id, lvl)] = counters.get((num_id, lvl), 0) + 1
            for k in [k for k in counters if k[0] == num_id and k[1] > lvl]:
                del counters[k]
            prefix = ".".join(str(counters.get((num_id, l), 1)) for l in range(lvl + 1))
            if not re.match(r"^\d", text):
                text = f"{prefix}. {text}"
        out.append(text)
    for table in d.tables:                       # таблицы оргструктуры
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(dict.fromkeys(cells)))
    return out


def load_docx(path: str, doc_id: str) -> tuple[Document, list[dict]]:
    paras = _docx_paragraph_texts(path)
    text = "\n".join(paras)
    meta = _meta_from_text(text[:3000])
    # У DOCX нет страниц: условная «страница» = блок из 40 абзацев
    pages = [(i // 40 + 1, "\n".join(paras[i:i + 40])) for i in range(0, len(paras), 40)]
    clauses, anomalies = split_into_clauses(pages, doc_id, meta["edition"])
    doc = Document(doc_id=doc_id, file=Path(path).name, kind="docx", clauses=clauses, **meta)
    return doc, anomalies


def load_xlsx(path: str, doc_id: str, edition: str = "") -> tuple[Document, list[dict]]:
    """Каждая непустая строка листа — отдельный «пункт» с адресом Лист!A{row}."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    clauses: list[Clause] = []
    for ws in wb.worksheets:
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if not vals:
                continue
            number = f"{ws.title}!A{r_idx}"
            clauses.append(Clause(
                clause_id=f"{doc_id}:{number}", doc_id=doc_id, edition=edition,
                number=number, text=" | ".join(vals), page=1, section=ws.title,
            ))
    doc = Document(doc_id=doc_id, file=Path(path).name, kind="xlsx", edition=edition, clauses=clauses)
    return doc, []


def load_any(path: str, doc_id: str) -> tuple[Document, list[dict]]:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return load_pdf(path, doc_id)
    if ext == ".docx":
        return load_docx(path, doc_id)
    if ext in (".xlsx", ".xlsm"):
        return load_xlsx(path, doc_id)
    raise ValueError(f"Неподдерживаемый формат: {ext}")
