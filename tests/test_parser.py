from app.parser.loader import load_pdf
from tests.conftest import AFTER, BEFORE


def test_metadata():
    b, _ = load_pdf(BEFORE, "D8")
    a, _ = load_pdf(AFTER, "D9")
    assert b.edition == "8" and a.edition == "9"
    assert "13" in b.approved and "7" in a.approved


def test_clauses_and_pages():
    b, _ = load_pdf(BEFORE, "D8")
    a, _ = load_pdf(AFTER, "D9")
    assert b.get("5.6.2").page == 10
    assert "группы контроля качества" in b.get("5.6.2").text
    assert a.get("3.4.а").text.startswith("Департамент ИТ-аудита")
    assert a.get("3.4").page == 6


def test_glued_clauses_are_split():
    # «...Общества. 3.10.Работники могут...» — склеенные пункты разделяются
    b, _ = load_pdf(BEFORE, "D8")
    assert b.get("3.10").text.startswith("Работники могут")
    assert b.get("5.4.4.б").text.startswith("выявления рисков")


def test_toc_is_ignored():
    a, _ = load_pdf(AFTER, "D9")
    assert all("ОБЩИЕ ПОЛОЖЕНИЯ 1" not in c.text for c in a.clauses)


def test_docx_table_retains_position_in_document(tmp_path):
    from docx import Document
    from app.parser.loader import load_docx
    doc = Document()
    doc.add_paragraph('1. Общие положения')
    doc.add_paragraph('1.1. Первый пункт')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Текст таблицы первого пункта'
    doc.add_paragraph('1.2. Второй пункт')
    path = tmp_path / 'ordered.docx'
    doc.save(path)
    parsed, _ = load_docx(str(path), 'D')
    assert 'Текст таблицы' in parsed.get('1.1').text
    assert 'Текст таблицы' not in parsed.get('1.2').text
